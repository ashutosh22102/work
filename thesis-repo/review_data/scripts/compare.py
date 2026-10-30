import json, sys
from collections import Counter
from openpyxl import load_workbook

XL = "SLR_Search_and_Screening.xlsx"
inc = {r[0]: r for r in load_workbook(XL)["Included studies"].iter_rows(min_row=2, values_only=True)}


def dcat(d):
    d = d.lower()
    if "conceptual" in d or "descriptive" in d: return "Conceptual"
    if "formal model" in d: return "Formal model"
    if "panel" in d or "archival" in d or "secondary data (pls" in d: return "Panel/archival"
    if "case" in d or "qualitative" in d: return "Qualitative/case"
    if "mixed" in d or "multi-method" in d or "multi-phase" in d or "llm" in d: return "Mixed"
    return "Survey"


def ocat(o):
    o = o.strip()
    parts = [p.strip() for p in o.split(",")] if not o.startswith(("Other", "FP (")) else [o]
    def one(p):
        pl = p.lower()
        if pl.startswith("sca") or "sustain" in pl and "compet" in pl: return "SCA"
        if pl.startswith("ca") or pl.startswith("compet"): return "CA"
        if pl.startswith("fp"):
            return "FP"
        if pl.startswith("innov") or "innovation" in pl: return "INNOV"
        if pl.startswith("resil"): return "RES"
        if pl.startswith("bmi"): return "BMI"
        return "OTHER"
    cats = [one(p) for p in parts]
    for c in ["SCA", "CA", "FP", "INNOV", "RES", "BMI", "OTHER"]:
        if c in cats: return c


mine = {}
for l in open("extraction_before_adjudication.txt", encoding="utf-8"):
    if l.startswith("#") or not l.strip(): continue
    p = l.rstrip("\n").split("|")
    q = p[18]
    mine[p[0]] = dict(tier=inc[p[0]][2], design=dcat(p[1]), outcome=ocat(p[11]), outcome_raw=p[11], p1=p[15], p2=p[16],
                      s=p[8] == "x", z=p[9] == "x", t=p[10] == "x", p3_raw=p[17],
                      **{f"q{i+1}": int(q[i]) for i in range(7)})

rec = {c["id"]: c for c in json.load(open("recode_run1.json", encoding="utf-8"))["result"]["coded"]}
assert set(rec) == set(mine), set(rec) ^ set(mine)


def kappa(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n / n
    return po, (po - pe) / (1 - pe) if pe < 1 else 1.0


FIELDS = ["tier", "design", "outcome", "p1", "p2", "s", "z", "t", "q1", "q2", "q3", "q4", "q5", "q6"]
ids = sorted(mine)
stats = {}
for f in FIELDS:
    a = [mine[i][f] for i in ids]
    b = [rec[i][f] for i in ids]
    po, k = kappa(a, b)
    stats[f] = dict(agree=round(po, 3), kappa=round(k, 3), n_disagree=sum(x != y for x, y in zip(a, b)))
    print(f"{f:8s} agree {po:.2f} kappa {k:.2f} disagree {stats[f]['n_disagree']}")

dis = {}
for i in ids:
    d = {f: [mine[i][f], rec[i][f]] for f in FIELDS if mine[i][f] != rec[i][f]}
    if d: dis[i] = d
print("studies with any disagreement:", len(dis))
print("outcome raw->cat:", sorted(Counter((mine[i]["outcome_raw"], mine[i]["outcome"]) for i in ids).items()))
print("recode q7:", Counter(rec[i]["q7"] for i in ids))
json.dump(dict(stats=stats, disagreements=dis), open("compare.json", "w", encoding="utf-8"), indent=1, ensure_ascii=False)
