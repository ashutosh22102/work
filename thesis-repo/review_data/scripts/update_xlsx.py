import json, shutil, io, contextlib
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
with contextlib.redirect_stdout(io.StringIO()):
    import compare as C

XL = "SLR_Search_and_Screening.xlsx"
shutil.copy(XL, "SLR_before_adjudication.xlsx")
F = json.load(open("final.json", encoding="utf-8"))
cmp_ = json.load(open("compare.json", encoding="utf-8"))
ex = {l.split("|")[0]: l.rstrip("\n").split("|") for l in open("extraction.txt", encoding="utf-8") if not l.startswith("#") and l.strip()}
wb = load_workbook(XL)

# tier changes (R179) in Included studies and Screening log
for name in ("Included studies", "Screening log"):
    ws = wb[name]
    hdr = [c.value for c in ws[1]]
    tcol = hdr.index("Tier") + 1 if "Tier" in hdr else None
    for row in ws.iter_rows(min_row=2):
        i = row[0].value
        if i in F and tcol and row[tcol - 1].value in ("Core", "Supporting") and row[tcol - 1].value != F[i]["tier"]:
            print(name, i, row[tcol - 1].value, "->", F[i]["tier"]); row[tcol - 1].value = F[i]["tier"]

ws = wb["PRISMA counts"]
for row in ws.iter_rows(min_row=2):
    if row[0].value and "core (e-commerce" in row[0].value: row[1].value = sum(f["tier"] == "Core" for f in F.values())
    if row[0].value and "supporting (other" in row[0].value: row[1].value = sum(f["tier"] == "Supporting" for f in F.values())

ws = wb["Data extraction"]
ws.cell(1, 21).value = "Evidence for P3 (text)"
ws.cell(1, 22).value = "Limitations (Q7 evidence)"
ws.cell(1, 29).value = "Q7 limitations (yes/no/not assessable)"
ws.cell(1, 30).value = "Quality score Q1-Q6 (0-6)"
extra = ["Quality band (high 6 / medium 4-5 / low 0-3)", "Evidence source", "Full-text source", "P3 families (codebook)", "Changed in adjudication (original -> final)", "Note"]
for k, h in enumerate(extra):
    ws.cell(1, 32 + k).value = h
    ws.cell(1, 32 + k).font = Font(bold=True)
DES = {"Survey": "Survey", "Conceptual": "Conceptual", "Qualitative/case": "Qualitative/case", "Panel/archival": "Panel/archival", "Mixed": "Mixed methods", "Formal model": "Formal model"}
Q7 = {"yes": "yes", "no": "no", "not_assessable": "not assessable"}
for row in ws.iter_rows(min_row=2):
    i = row[0].value
    f, p = F[i], ex[i]
    row[4].value = f["tier"]
    row[7].value = p[1]
    if i == "R041": row[8].value = p[2]
    row[11].value, row[12].value, row[13].value = p[8], p[9], p[10]
    row[14].value = p[11]
    row[15].value = p[12]
    row[18].value, row[19].value = f["p1"], f["p2"]
    row[21].value = f["ft_note"] if f["evidence"] == "Full text" else "Abstract only; limitations not assessable unless stated in abstract"
    for n in range(6): row[22 + n].value = f[f"q{n+1}"]
    row[28].value = Q7[f["q7"]]
    row[29].value = f["q16"]
    vals = [f["band"], "Full text verified" if f["evidence"] == "Full text" else "Abstract-based", f["ft_url"], ", ".join(f["p3"]),
            "; ".join(f"{k}: {a} -> {b}" for k, (a, b) in f["changed"].items()), f["note"]]
    for k, v in enumerate(vals):
        ws.cell(row[0].row, 32 + k).value = v

# second coding sheet
if "Second coding" in wb.sheetnames: del wb["Second coding"]
ws = wb.create_sheet("Second coding")
FL = C.FIELDS
hdr = ["ID"] + [f"A {k}" for k in FL] + [f"B {k}" for k in FL] + ["B q7", "B rationale"] + [f"Final {k}" for k in FL] + ["Final q7", "Evidence source"]
ws.append(hdr)
for i in sorted(F):
    a, b, f = C.mine[i], C.rec[i], F[i]
    ws.append([i] + [str(a[k]) for k in FL] + [str(b[k]) for k in FL] + [b["q7"], b["rationale"]] + [str(f[k]) for k in FL] + [f["q7"], f["evidence"]])
ws.append([])
ws.append(["Agreement between coding A (original) and coding B (independent, codebook-based, AI-assisted)"])
ws.append(["Field", "Percent agreement", "Cohen's kappa", "Disagreements"])
for k, v in cmp_["stats"].items():
    ws.append([k, v["agree"], v["kappa"], v["n_disagree"]])
ws.append(["Note: disagreements were resolved by re-applying the codebook (codebook.md); kappa is unstable for items with very skewed distributions (e.g. q1)."])
for c in ws[1]: c.font = Font(bold=True)
wb.save(XL)
print("saved")
