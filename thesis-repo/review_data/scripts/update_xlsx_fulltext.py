"""Write the full-text verified codes (fulltext_verified.json, already applied to final.json) into the Excel workbook.

Updates the 'Data extraction' and 'Second coding' (Final columns) sheets for every accessed study.
Backs up the workbook to SLR_before_fulltext.xlsx first. Run from review_data/ after apply_fulltext.py.
"""
import json, shutil

from openpyxl import load_workbook

XL = "SLR_Search_and_Screening.xlsx"
shutil.copy(XL, "SLR_before_fulltext.xlsx")
F = json.load(open("final.json", encoding="utf-8"))
V = json.load(open("fulltext_verified.json", encoding="utf-8"))
ex = {l.split("|")[0]: l.rstrip("\r\n").split("|") for l in open("extraction.txt", encoding="utf-8") if not l.startswith("#") and l.strip()}
acc = [i for i, v in V.items() if v.get("accessed")]
wb = load_workbook(XL)
Q7 = {"yes": "yes", "no": "no", "not_assessable": "not assessable"}

ws = wb["Data extraction"]
for row in ws.iter_rows(min_row=2):
    i = row[0].value
    if i not in acc:
        continue
    f, p, v = F[i], ex[i], V[i]
    row[11].value, row[12].value, row[13].value = p[8], p[9], p[10]
    row[14].value = p[11]
    row[18].value, row[19].value = f["p1"], f["p2"]
    row[21].value = v["limitations_evidence"]
    for n in range(6):
        row[22 + n].value = f[f"q{n+1}"]
    row[28].value = Q7[f["q7"]]
    row[29].value = f["q16"]
    row[31].value = f["band"]
    row[32].value = "Full text verified"
    row[33].value = v["source"]
    ft = "; ".join(f"{k[3:]}: {a} -> {b}" for k, (a, b) in f["changed"].items() if k.startswith("ft_"))
    row[36].value = "Full-text check: " + (ft or "abstract coding confirmed") + ". " + v["notes"]

ws = wb["Second coding"]
hdr = [c.value for c in ws[1]]
col = {h: k for k, h in enumerate(hdr)}
for row in ws.iter_rows(min_row=2):
    i = row[0].value
    if i not in acc:
        continue
    f = F[i]
    for k in ["tier", "design", "outcome", "p1", "p2", "s", "z", "t", "q1", "q2", "q3", "q4", "q5", "q6"]:
        row[col[f"Final {k}"]].value = str(f[k])
    row[col["Final q7"]].value = f["q7"]
    row[col["Evidence source"]].value = f["evidence"]

wb.save(XL)
print("updated", len(acc), "studies in", XL)
