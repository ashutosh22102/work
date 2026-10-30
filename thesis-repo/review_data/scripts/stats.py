import json
from collections import Counter
from openpyxl import load_workbook

F = json.load(open("final.json", encoding="utf-8"))
inc = {r[0]: r for r in load_workbook("SLR_Search_and_Screening.xlsx")["Included studies"].iter_rows(min_row=2, values_only=True)}
vhb = {i for i in inc if inc[i][11] == "yes"}
ids = sorted(F)
core = [i for i in ids if F[i]["tier"] == "Core"]
C_ = lambda xs: dict(Counter(xs))
print("tier", C_(F[i]["tier"] for i in ids), "core:", core)
print("design", C_(F[i]["design"] for i in ids))
q = [F[i]["q16"] for i in ids]
print("band", C_(F[i]["band"] for i in ids), "mean q1-6 %.2f" % (sum(q) / len(q)))
print("vhb high", sum(F[i]["band"] == "high" for i in vhb), "of", len(vhb), "| nonvhb high", sum(F[i]["band"] == "high" for i in ids if i not in vhb), "of", len(ids) - len(vhb))
print("q7", C_(F[i]["q7"] for i in ids), "evidence", C_(F[i]["evidence"] for i in ids), "q7 by evidence", C_((F[i]["evidence"], F[i]["q7"]) for i in ids))
print("q7 no:", [i for i in ids if F[i]["q7"] == "no"])
for n in range(1, 7): print(f"q{n} met", sum(F[i][f"q{n}"] for i in ids))
print("outcome x tier", sorted(C_((F[i]["outcome"], F[i]["tier"]) for i in ids).items()))
for o in ["FP", "CA", "SCA", "INNOV", "RES", "BMI", "OTHER"]:
    s = [i for i in ids if F[i]["outcome"] == o]
    print(o, len(s), "p1:", C_(F[i]["p1"] for i in s), "core:", [i for i in s if F[i]["tier"] == "Core"])
    if o in ("SCA",): print("   SCA ids", [(i, F[i]["p1"], F[i]["design"]) for i in s])
p1 = [i for i in ids if F[i]["p1"] != "na"]
print("P1 relevant", len(p1), C_(F[i]["p1"] for i in p1))
print("P1 by value ids", {v: [i for i in p1 if F[i]["p1"] == v] for v in ["0", "c", "-"]})
print("P1 + outside SCA/CA/FP/INNOV", [i for i in p1 if F[i]["outcome"] not in ("SCA", "CA", "FP", "INNOV")])
hq = [i for i in ids if F[i]["band"] == "high"]
print("P1 high", len([i for i in p1 if i in hq]), sum(F[i]["p1"] == "+" for i in p1 if i in hq))
pc = [i for i in core if F[i]["p1"] != "na"]
print("P1 core", len(pc), C_(F[i]["p1"] for i in pc))
p2 = [i for i in ids if F[i]["p2"] != "na"]
print("P2 relevant", len(p2), C_(F[i]["p2"] for i in p2), {v: [i for i in p2 if F[i]["p2"] == v] for v in ["0", "c"]})
print("P2 high", len([i for i in p2 if i in hq]), sum(F[i]["p2"] == "+" for i in p2 if i in hq))
print("P2 core", len([i for i in core if F[i]["p2"] != "na"]), C_(F[i]["p2"] for i in core if F[i]["p2"] != "na"))
print("S Z T", sum(F[i]["s"] for i in ids), sum(F[i]["z"] for i in ids), sum(F[i]["t"] for i in ids), "all3", sum(F[i]["s"] and F[i]["z"] and F[i]["t"] for i in ids))
print("p3 families", Counter(x for i in ids for x in F[i]["p3"]).most_common(), "studies with any", sum(bool(F[i]["p3"]) for i in ids))
emp = [i for i in ids if F[i]["design"] != "Conceptual"]
def sub(s, name):
    a = [i for i in s if F[i]["p1"] != "na"]; b = [i for i in s if F[i]["p2"] != "na"]
    print(name, len(s), "P1", sum(F[i]["p1"] == "+" for i in a), "/", len(a), C_(F[i]["p1"] for i in a), "| P2", sum(F[i]["p2"] == "+" for i in b), "/", len(b))
sub(emp, "empirical"); sub([i for i in ids if F[i]["band"] != "low"], "no-low"); sub([i for i in ids if i in vhb], "vhb"); sub(hq, "high")
print("core high", [i for i in core if F[i]["band"] == "high"], "core vhb", sum(i in vhb for i in core))
print("core bands", C_(F[i]["band"] for i in core))
