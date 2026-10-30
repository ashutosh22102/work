"""Write fulltext_todo.md: every study still coded from its abstract, core studies first."""
import json
from openpyxl import load_workbook

F = json.load(open("final.json", encoding="utf-8"))
inc = {r[0]: r for r in load_workbook("SLR_Search_and_Screening.xlsx")["Included studies"].iter_rows(min_row=2, values_only=True)}
todo = sorted((i for i in F if F[i]["evidence"] != "Full text"), key=lambda i: (F[i]["tier"] != "Core", i))
L = ["# Full-text verification checklist", "",
     f"{len(todo)} studies are still coded from their abstracts only. Core studies come first. For each one, read the full text and",
     "check outcome, P1, P2, S/Z/T, sample, Q1-Q6 and Q7 against `codebook.md`. Record the results in `fulltext_verified.json` (see CLAUDE.md, step 2).", "",
     "| # | ID | Tier | First author (year) | Title | DOI | Current: design / outcome / P1 / P2 / Q1-Q6 / Q7 |", "|---|---|---|---|---|---|---|"]
for n, i in enumerate(todo, 1):
    r, f = inc[i], F[i]
    title = (r[5] or "").replace("|", "/").replace("\n", " ")[:90]
    L.append(f"| {n} | {i} | {f['tier']} | {r[3]} ({r[4]}) | {title} | {r[10] or '-'} | {f['design']} / {f['outcome']} / {f['p1']} / {f['p2']} / {f['q16']} / {f['q7']} |")
open("fulltext_todo.md", "w", encoding="utf-8").write("\n".join(L) + "\n")
print(len(todo), "core:", sum(F[i]["tier"] == "Core" for i in todo))
